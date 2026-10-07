# -*- coding: utf-8 -*-
"""
ledger.jsonl を唯一の正として、現金出納帳.xlsx を月別シート構成で
毎回ゼロから組み立て直す。外部のテンプレートファイルには依存しない
（その場で罫線・フォント・列幅を設定するため、テンプレートの有無や
　状態に左右されずいつでも再現できる）。
"""
import json
import math
import unicodedata
from copy import copy
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

COMPANY_NAME = "株式会社セリエコンチェルト"

THIN = Side(style="thin")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT_NUM = Font(name="Yu Gothic", size=10)
FONT_TEXT = Font(name="Noto Sans CJK SC", size=10)
FONT_TOTAL = Font(name="Yu Gothic", size=10, bold=True)
FONT_TOTAL_TEXT = Font(name="Noto Sans CJK SC", size=10, bold=True)
FILL_TOTAL = PatternFill(fill_type="solid", fgColor="FFDDEBF7")
NUMFMT_YEN = "#,##0;(#,##0);-"

COL_WIDTHS = {"A": 12, "B": 16, "C": 52, "D": 15, "E": 15, "F": 16}

CAPACITY_PER_LINE = 48
ROW_HEIGHT_PER_LINE = 16
MIN_ROW_HEIGHT = 18
ROW_HEIGHT_PADDING = 6
EXTRA_LINE_SAFETY = 1


def load_ledger(ledger_path):
    rows = []
    with open(ledger_path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                rows.append(json.loads(line))
    rows.sort(key=lambda r: (r["date"], r.get("file") or ""))
    return rows


def build_summary(r):
    if r.get("summary_override"):
        summary = r["summary_override"]
    else:
        summary = r.get("store") or ""
        if r.get("items"):
            summary += f"\u3000{r['items']}"
        if r.get("reg"):
            summary += f"(登録番号 {r['reg']})"
    if r.get("status") == "needs_review":
        tag = "[自動判定・要確認]" if r.get("source", "").startswith("line_auto") else "[要確認]"
        note = r.get("note") or ""
        summary += f"{tag}{(': ' + note) if note else ''}"
    return summary


def _effective_width(text):
    total = 0
    for ch in str(text):
        total += 2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1
    return total


def _row_height(text):
    if not text:
        return MIN_ROW_HEIGHT
    lines = max(1, math.ceil(_effective_width(text) / CAPACITY_PER_LINE)) + EXTRA_LINE_SAFETY
    return max(MIN_ROW_HEIGHT, lines * ROW_HEIGHT_PER_LINE + ROW_HEIGHT_PADDING)


def _style_data_row(ws, row):
    for col in range(1, 7):
        cell = ws.cell(row=row, column=col)
        cell.border = BORDER_ALL
        cell.font = FONT_TEXT if col in (2, 3) else FONT_NUM
        wrap = col in (2, 3)
        halign = "center" if col == 1 else ("right" if col in (4, 5, 6) else "left")
        cell.alignment = Alignment(horizontal=halign, vertical="center", wrap_text=wrap)
        if col in (4, 5, 6):
            cell.number_format = NUMFMT_YEN


def _style_total_row(ws, row):
    for col in range(1, 7):
        cell = ws.cell(row=row, column=col)
        cell.border = BORDER_ALL
        cell.fill = FILL_TOTAL
        cell.font = FONT_TOTAL_TEXT if col == 3 else FONT_TOTAL
        halign = "left" if col == 3 else ("right" if col in (4, 5, 6) else "general")
        cell.alignment = Alignment(horizontal=halign, vertical="center")
        if col in (4, 5, 6):
            cell.number_format = NUMFMT_YEN


def build_sheet(wb, title, month_label, entries, carry_in):
    ws = wb.create_sheet(title=title)
    for col, width in COL_WIDTHS.items():
        ws.column_dimensions[col].width = width

    ws.merge_cells("A1:F1")
    ws["A1"] = "現金出納帳"
    ws.merge_cells("A2:F2")
    ws["A2"] = COMPANY_NAME
    ws["A3"] = "年月："
    ws["B3"] = month_label

    ws["A5"] = "日付"
    ws["B5"] = "勘定科目"
    ws["C5"] = "摘要"
    ws["D5"] = "入金（収入）"
    ws["E5"] = "出金（支出）"
    ws["F5"] = "差引残高"
    for col in range(1, 7):
        c = ws.cell(row=5, column=col)
        c.font = FONT_TOTAL
        c.border = BORDER_ALL
        c.alignment = Alignment(horizontal="center", vertical="center")

    ws["C6"] = "前月繰越"
    ws["F6"] = carry_in
    for col in range(1, 7):
        ws.cell(row=6, column=col).border = BORDER_ALL
        ws.cell(row=6, column=col).font = FONT_NUM
    ws.cell(row=6, column=6).number_format = NUMFMT_YEN
    ws.row_dimensions[6].height = 18

    row = 7
    income_total = 0
    expense_total = 0
    for e in entries:
        dt = datetime.strptime(e["date"], "%Y-%m-%d")
        ws.cell(row=row, column=1).value = dt
        ws.cell(row=row, column=1).number_format = "m/d"
        ws.cell(row=row, column=2).value = e.get("category") or ""
        summary = build_summary(e)
        ws.cell(row=row, column=3).value = summary
        if e.get("is_income"):
            ws.cell(row=row, column=4).value = e["amount"]
            income_total += e["amount"]
        else:
            ws.cell(row=row, column=5).value = e["amount"]
            expense_total += e["amount"]
        ws.cell(row=row, column=6).value = f"=F{row - 1}+N(D{row})-N(E{row})"
        _style_data_row(ws, row)
        ws.row_dimensions[row].height = _row_height(summary)
        row += 1

    total_row = row
    ws.cell(row=total_row, column=3).value = "合計"
    ws.cell(row=total_row, column=4).value = f"=SUM(D7:D{total_row - 1})" if total_row > 7 else 0
    ws.cell(row=total_row, column=5).value = f"=SUM(E7:E{total_row - 1})" if total_row > 7 else 0
    ws.cell(row=total_row, column=6).value = f"=F{total_row - 1}" if total_row > 7 else carry_in
    _style_total_row(ws, total_row)

    carry_out = carry_in + income_total - expense_total
    return carry_out


def build_workbook(ledger_path, output_paths):
    entries = load_ledger(ledger_path)
    by_month = {}
    for e in entries:
        by_month.setdefault(e["date"][:7], []).append(e)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    carry = 0
    for month in sorted(by_month):
        y, m = month.split("-")
        label = f"{y}年 {int(m)}月"
        title = f"{y}年{int(m)}月"
        carry = build_sheet(wb, title, label, by_month[month], carry)

    saved = []
    for path in output_paths:
        wb.save(path)
        saved.append(path)
    return saved, carry


if __name__ == "__main__":
    import os
    LEDGER = os.path.join(os.path.dirname(__file__), "ledger.jsonl")
    OUT1 = r"C:\Users\user\Desktop\領収書\現金出納帳_セリエコンチェルト (1).xlsx"
    OUT2 = r"C:\Users\user\Downloads\現金出納帳_セリエコンチェルト (1).xlsx"
    saved, final_balance = build_workbook(LEDGER, [OUT1, OUT2])
    print("saved:", saved)
    print("final balance:", final_balance)
